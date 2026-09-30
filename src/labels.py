"""Hand-label sample for validating the LLM attribute extraction.

    python -m src.labels create   ->  data/labels/hand_labels.csv (never overwrites)
    python -m src.labels xlsx     ->  data/local/hand_labels_with_text.xlsx to label in Excel
                                      (built from the CSV; never overwrites)
    python -m src.labels import   ->  checks the finished xlsx and writes its labels
                                      into hand_labels.csv (only if there are no problems)

The shops' card text and descriptions are shown to the labeller but never
committed: the CSV holds IDs, titles and labels; the workbook, which adds the
text from data/local/product_text.csv, lives in git-ignored data/local/.
(data/labels/hand_labels.xlsx is the finished workbook with that text removed.)

100 products, stratified: at least MIN_PER_ROASTER from every roaster, the rest
in proportion to roaster size. Within a roaster, Phase 2's ambiguous products
(is_coffee unresolved) are taken first, up to half its quota, then the rest is
split between products with and without a live description. Fixed seed.

The file shows only the text the model will see, never any pipeline label,
so labelling stays blind. Note lines at the top start with "#"; read_labels()
skips them.
"""
import argparse
import math

import pandas as pd

from src import config
from src.product_inputs import OUT as INPUTS
from src.product_inputs import SHOP_TEXT, with_text

LABELS_DIR = config.DATA_DIR / "labels"
LABELS = LABELS_DIR / "hand_labels.csv"
SAMPLE_SIZE = 100
MIN_PER_ROASTER = 5
SEED = 42

# Product columns in the committed CSV. The shops' own text (card_text,
# description) is added only in the local workbook: see create_xlsx.
TEXT_COLUMNS = ["roaster", "product_id", "product_title", "other_titles", "product_type",
                "variant_titles"]
# The fields to label, with the allowed values (same as the extraction schema).
FIELDS = {
    "is_coffee": "yes / no / unknown",
    "origin_country": "country name, 'multiple' or unknown",
    "region": "e.g. Chikmagalur, Coorg, Araku Valley; or unknown",
    "estate": "estate or farm name; or unknown",
    "species": "arabica / robusta / blend / other / unknown",
    "origin_type": "single_origin / blend / unknown",
    "process": "washed / natural / honey / monsooned / experimental / mixed / unknown",
    "roast_level": "light / light_medium / medium / medium_dark / dark / omni / unknown",
    "decaf": "yes / no / unknown",
    "flavoured_infused": "yes / no / unknown",
}
# Allowed values per field: a list = dropdown (and the only values accepted
# on import); None = free text ("unknown" when the text doesn't say).
ALLOWED: dict[str, list[str] | None] = {
    "is_coffee": ["yes", "no", "unknown"],
    "origin_country": None,
    "region": None,
    "estate": None,
    "species": ["arabica", "robusta", "blend", "other", "unknown"],
    "origin_type": ["single_origin", "blend", "unknown"],
    "process": ["washed", "natural", "honey", "monsooned", "experimental", "mixed", "unknown"],
    "roast_level": ["light", "light_medium", "medium", "medium_dark", "dark", "omni", "unknown"],
    "decaf": ["yes", "no", "unknown"],
    "flavoured_infused": ["yes", "no", "unknown"],
}
assert list(ALLOWED) == list(FIELDS)
XLSX = config.DATA_DIR / "local" / "hand_labels_with_text.xlsx"   # git-ignored: has shop text
NOTE = [
    "# LABEL ONLY FROM THE TEXT PROVIDED IN THIS ROW. Use \"unknown\" when the text "
    "doesn't say. Don't look anything up or use outside knowledge.",
    "# Allowed values -> " + " | ".join(f"{k}: {v}" for k, v in FIELDS.items()),
    "# species 'other' = excelsa, liberica etc. process 'experimental' = anaerobic, "
    "carbonic maceration, yeast/koji/culture fermentation. Save as CSV UTF-8.",
]


def allocate(sizes: pd.Series, total: int = SAMPLE_SIZE, floor: int = MIN_PER_ROASTER) -> dict:
    """Per-roaster quotas: `floor` each (or all, if a roaster has fewer), the
    remainder split in proportion to how many products remain above the floor.
    Largest remainders get the leftover units so quotas sum exactly to `total`."""
    base = {r: min(floor, n) for r, n in sizes.items()}
    spare = {r: sizes[r] - base[r] for r in sizes.index}
    left = total - sum(base.values())
    raw = {r: left * spare[r] / sum(spare.values()) for r in sizes.index}
    quota = {r: base[r] + math.floor(raw[r]) for r in sizes.index}
    for r in sorted(raw, key=lambda r: raw[r] - math.floor(raw[r]), reverse=True):
        if sum(quota.values()) >= total:
            break
        if quota[r] < sizes[r]:
            quota[r] += 1
    return quota


def sample_roaster(g: pd.DataFrame, n: int, seed: int) -> pd.DataFrame:
    """Ambiguous products first (up to half of n), then a description / no-description mix."""
    is_amb = g.is_coffee_guess_p2.isna() | (g.is_coffee_guess_p2 == "")
    take_amb = g[is_amb].sample(min(is_amb.sum(), n // 2), random_state=seed)
    # The rest of the quota comes from NON-ambiguous products, so the "half"
    # cap holds; leftover ambiguous ones only fill in if a roaster runs short.
    rest = g[~is_amb]
    k = n - len(take_amb)
    with_desc = rest[rest.has_description]
    n_desc = min(len(with_desc), round(k * len(with_desc) / len(rest))) if len(rest) else 0
    picked = [take_amb, with_desc.sample(n_desc, random_state=seed)]
    remaining = rest.drop(picked[1].index)
    picked.append(remaining.sample(min(k - n_desc, len(remaining)), random_state=seed))
    short = n - sum(len(p) for p in picked)
    if short > 0:
        spare_amb = g[is_amb].drop(take_amb.index)
        picked.append(spare_amb.sample(min(short, len(spare_amb)), random_state=seed))
    return pd.concat(picked)


def create() -> None:
    if LABELS.exists():
        raise SystemExit(f"{LABELS} already exists; not overwriting hand labels.")
    inputs = pd.read_csv(INPUTS, dtype={"product_id": str, "is_coffee_guess_p2": str})
    quota = allocate(inputs.groupby("roaster").size())
    sample = pd.concat(sample_roaster(g, quota[r], SEED) for r, g in inputs.groupby("roaster"))
    # Shuffle so roasters are interleaved (less anchoring on one shop's style).
    sample = sample.sample(frac=1, random_state=SEED).reset_index(drop=True)
    out = sample[TEXT_COLUMNS].fillna("")
    out.insert(0, "label_id", range(1, len(out) + 1))
    for field in FIELDS:
        out[field] = ""
    out["labeller_notes"] = ""
    LABELS_DIR.mkdir(parents=True, exist_ok=True)
    # utf-8-sig so Excel shows ₹ and accented names correctly.
    with LABELS.open("w", encoding="utf-8-sig", newline="") as f:
        f.write("\n".join(NOTE) + "\n")
        out.to_csv(f, index=False)
    print(f"Wrote {len(out)} products to {LABELS}")
    print(sample.groupby("roaster").size().to_string())
    amb = sample.is_coffee_guess_p2.isna() | (sample.is_coffee_guess_p2 == "")
    print(f"ambiguous: {amb.sum()}, with description: {sample.has_description.sum()}, "
          f"title/card only: {(~sample.has_description).sum()}")


def read_labels(path=None) -> pd.DataFrame:
    """The labels CSV, skipping the '#' note lines at the top. All text (dtype=str),
    so product IDs stay exact."""
    path = path or LABELS
    with path.open(encoding="utf-8-sig") as f:
        skip = 0
        for line in f:
            if not line.startswith("#"):
                break
            skip += 1
    return pd.read_csv(path, skiprows=skip, dtype=str, encoding="utf-8-sig").fillna("")


# ---------- Excel workbook for labelling ----------

# Column widths (Excel character units) and which columns wrap text.
WIDTHS = {"label_id": 6, "roaster": 14, "product_id": 16, "product_title": 34,
          "other_titles": 28, "product_type": 14, "variant_titles": 40, "card_text": 40,
          "description": 70, "labeller_notes": 30}
WRAP = {"product_title", "other_titles", "variant_titles", "card_text", "description",
        "labeller_notes"}
FROZEN = "E2"   # header row + label_id, roaster, product_id, product_title stay visible
INSTRUCTIONS = [
    ("RULE", "Label ONLY from the text in the row. Use \"unknown\" when the text doesn't say. "
             "Don't look anything up or use outside knowledge."),
    ("HOW", "Fill the grey label columns on the Labels sheet. Dropdown columns only accept "
            "the listed values. Don't edit the product columns. Save as .xlsx (not CSV)."),
    ("species: other", "excelsa, liberica etc."),
    ("process: experimental", "anaerobic, carbonic maceration, yeast / koji / culture fermentation"),
    ("process: mixed", "the text names more than one process for the same coffee"),
    ("origin_type: blend", "several origins or estates combined (not arabica+robusta alone: "
                           "that's species = blend)"),
]


def create_xlsx(csv_path=None, xlsx_path=None) -> None:
    """Build the labelling workbook from hand_labels.csv (which isn't changed)."""
    from openpyxl import Workbook
    from openpyxl.styles import Alignment, Font, PatternFill
    from openpyxl.utils import get_column_letter
    from openpyxl.worksheet.datavalidation import DataValidation

    csv_path, xlsx_path = csv_path or LABELS, xlsx_path or XLSX
    if xlsx_path.exists():
        raise SystemExit(f"{xlsx_path} already exists; not overwriting your labels.")
    df = read_labels(csv_path)
    # The labeller sees the shops' text too, placed just before the label columns.
    df = with_text(df)
    cols = [c for c in df.columns if c not in SHOP_TEXT]
    at = cols.index(next(iter(FIELDS)))
    df = df[cols[:at] + SHOP_TEXT + cols[at:]]
    wb = Workbook()

    # Instructions sheet.
    ins = wb.active
    ins.title = "Instructions"
    ins.append(["Topic", "Guidance"])
    for row in INSTRUCTIONS:
        ins.append(list(row))
    ins.append([])
    ins.append(["Field", "Allowed values"])
    for field, allowed in ALLOWED.items():
        ins.append([field, " / ".join(allowed) if allowed else FIELDS[field] + " (free text)"])
    ins.column_dimensions["A"].width, ins.column_dimensions["B"].width = 24, 110
    for row in ins.iter_rows():
        row[1].alignment = Alignment(wrap_text=True, vertical="top")
    for cell in (ins["A1"], ins["B1"], ins[f"A{len(INSTRUCTIONS) + 3}"], ins[f"B{len(INSTRUCTIONS) + 3}"]):
        cell.font = Font(bold=True)

    # Labels sheet.
    ws = wb.create_sheet("Labels")
    cols = list(df.columns)
    ws.append(cols)
    for record in df.itertuples(index=False):
        ws.append([str(v) for v in record])
    label_fill = PatternFill("solid", fgColor="EDEDED")
    n = len(df) + 1
    for j, col in enumerate(cols, start=1):
        letter = get_column_letter(j)
        ws.column_dimensions[letter].width = WIDTHS.get(col, 16)
        ws[f"{letter}1"].font = Font(bold=True)
        for i in range(2, n + 1):
            cell = ws[f"{letter}{i}"]
            if col == "product_id":
                cell.number_format = "@"    # text: Excel can't turn it into 7.31E+12
            cell.alignment = Alignment(wrap_text=col in WRAP, vertical="top")
            if col in ALLOWED or col == "labeller_notes":
                cell.fill = label_fill
                cell.value = None          # blank, not the string "" from the CSV
        allowed = ALLOWED.get(col)
        if allowed:
            dv = DataValidation(type="list", formula1='"' + ",".join(allowed) + '"',
                                allow_blank=True, showErrorMessage=True,
                                errorTitle="Not an allowed value",
                                error="Pick one of: " + ", ".join(allowed))
            dv.add(f"{letter}2:{letter}{n}")
            ws.add_data_validation(dv)
    ws.freeze_panes = FROZEN
    wb.active = wb.sheetnames.index("Labels")
    xlsx_path.parent.mkdir(parents=True, exist_ok=True)
    wb.save(xlsx_path)
    print(f"Wrote {len(df)} products to {xlsx_path}")


def _cell_text(value) -> str:
    """Excel cell -> clean string. Whole numbers lose any '.0' (a product ID
    Excel stored as a number); None -> ''."""
    if value is None:
        return ""
    if isinstance(value, float) and value.is_integer():
        return str(int(value))
    return str(value).strip()


def check_labels(labelled: pd.DataFrame, original: pd.DataFrame) -> list[str]:
    """Problems with the labelled sheet, as readable messages (empty = all good)."""
    problems = []
    missing_cols = [c for c in original.columns if c not in labelled.columns]
    if missing_cols:
        return [f"missing columns: {missing_cols}"]
    if len(labelled) != len(original):
        problems.append(f"{len(labelled)} rows, expected {len(original)}")
    orig = original.set_index("label_id")
    for r in labelled.itertuples(index=False):
        row = r._asdict()
        lid = row["label_id"]
        if lid not in orig.index:
            problems.append(f"label_id {lid}: not in the original sample")
            continue
        for key in ("product_id", "roaster"):
            if row[key] != orig.loc[lid, key]:
                problems.append(f"label_id {lid}: {key} is {row[key]!r}, expected {orig.loc[lid, key]!r}")
        for field, allowed in ALLOWED.items():
            value = row[field].strip()
            if not value:
                problems.append(f"label_id {lid}: {field} is blank")
            elif allowed and value.lower() not in allowed:
                problems.append(f"label_id {lid}: {field} = {value!r} not in {allowed}")
    return problems


def import_xlsx(csv_path=None, xlsx_path=None) -> list[str]:
    """Check the finished workbook; if clean, write its labels into hand_labels.csv.

    The CSV's '#' note lines are kept, and the previous CSV is saved as
    hand_labels.csv.bak first. Returns the list of problems (empty on success).
    """
    from openpyxl import load_workbook

    csv_path, xlsx_path = csv_path or LABELS, xlsx_path or XLSX
    original = read_labels(csv_path)
    ws = load_workbook(xlsx_path, read_only=True, data_only=True)["Labels"]
    rows = list(ws.iter_rows(values_only=True))
    header = [_cell_text(h) for h in rows[0]]
    labelled = pd.DataFrame([[_cell_text(v) for v in r] for r in rows[1:] if any(v is not None for v in r)],
                            columns=header)
    problems = check_labels(labelled, original)
    if problems:
        print(f"{len(problems)} problem(s); hand_labels.csv NOT changed:")
        for p in problems[:50]:
            print("  -", p)
        return problems
    for field in ALLOWED:   # dropdown values stored lower-case, free text as typed
        if ALLOWED[field]:
            labelled[field] = labelled[field].str.lower()
    with csv_path.open(encoding="utf-8-sig") as f:
        notes = [line for line in f if line.startswith("#")]
    csv_path.with_suffix(".csv.bak").write_bytes(csv_path.read_bytes())
    with csv_path.open("w", encoding="utf-8-sig", newline="") as f:
        f.write("".join(notes))
        labelled[list(original.columns)].to_csv(f, index=False)
    print(f"All {len(labelled)} rows valid; labels written to {csv_path} (backup: .csv.bak)")
    return []


# ---------- labelling rule: Indian region -> origin_country India ----------

# Every Indian place used as a region in the labels (lower case). Decided
# 2026-09-30 and applied to ALL 100 labels (not tuned to model output):
# where region is one of these, origin_country = India.
INDIAN_REGIONS = {
    "araku valley", "assam", "baba budan giri", "biligirirangan hills", "chikkamagaluru",
    "chikmagalur", "coorg", "gandha", "karnataka", "khasi hills", "kohima",
    "kolasib, mizoram", "malabar", "nilgiri", "nilgiris", "paderu", "sakleshpur",
    "thandigudi", "valparai", "zora, mizoram",
}


def india_rule_changes(df: pd.DataFrame) -> tuple[list[int], list[str]]:
    """(row positions to set to India, conflicts). Only 'unknown' is ever filled;
    a different country already recorded is reported, not overwritten."""
    change, conflicts = [], []
    for i, r in enumerate(df.itertuples(index=False)):
        if str(r.region).strip().lower() not in INDIAN_REGIONS:
            continue
        country = str(r.origin_country).strip().lower()
        if country == "unknown":
            change.append(i)
        elif country != "india":
            conflicts.append(f"label_id {r.label_id}: region {r.region!r} but country {r.origin_country!r}")
    return change, conflicts


def apply_india_rule(xlsx_path=None, csv_path=None) -> None:
    """Apply the rule in the workbook (the labelling source), then re-import to the CSV."""
    from openpyxl import load_workbook

    xlsx_path, csv_path = xlsx_path or XLSX, csv_path or LABELS
    wb = load_workbook(xlsx_path)
    ws = wb["Labels"]
    header = [c.value for c in ws[1]]
    rows = [[_cell_text(c.value) for c in r] for r in ws.iter_rows(min_row=2)]
    df = pd.DataFrame([r for r in rows if any(r)], columns=header)
    change, conflicts = india_rule_changes(df)
    for c in conflicts:
        print("  CONFLICT (not changed):", c)
    col = header.index("origin_country") + 1
    xlsx_path.with_suffix(".xlsx.bak").write_bytes(xlsx_path.read_bytes())
    for i in change:
        ws.cell(row=i + 2, column=col, value="India")
        r = df.iloc[i]
        print(f"  label_id {r.label_id:>3} | {r.roaster:17} | region {r.region!r:20} | "
              f"{r.product_title[:50]}")
    wb.save(xlsx_path)
    print(f"Set origin_country = India on {len(change)} rows (backup: .xlsx.bak)")
    if import_xlsx(csv_path, xlsx_path):
        raise SystemExit(1)


def main() -> None:
    parser = argparse.ArgumentParser(description="Hand-label sample")
    parser.add_argument("action", choices=["create", "xlsx", "import", "india-rule"])
    args = parser.parse_args()
    if args.action == "create":
        create()
    elif args.action == "xlsx":
        create_xlsx()
    elif args.action == "india-rule":
        apply_india_rule()
    elif import_xlsx():
        raise SystemExit(1)


if __name__ == "__main__":
    main()
