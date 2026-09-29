"""Hand-label sampling and reading."""
import pandas as pd

from src import labels


def test_allocate_floor_and_total():
    sizes = pd.Series({"A": 87, "B": 10, "C": 16, "D": 3})
    q = labels.allocate(sizes, total=30, floor=5)
    assert sum(q.values()) == 30
    assert q["D"] == 3                      # can't take more than exist
    assert q["B"] >= 5 and q["C"] >= 5
    assert q["A"] > q["C"] > q["B"] - 1     # the rest goes in proportion to size


def test_allocate_real_shape_sums_to_100():
    sizes = pd.Series({"Araku": 16, "Black Baza": 45, "Bloom": 52, "Blue Tokai": 87,
                       "Corridor Seven": 78, "Devans": 62, "Dope Coffee": 25, "Grey Soul": 63,
                       "Kapi Kottai": 36, "Subko": 62, "Third Wave Coffee": 10})
    q = labels.allocate(sizes)
    assert sum(q.values()) == 100 and min(q.values()) >= 5


def test_sample_roaster_takes_ambiguous_first():
    g = pd.DataFrame({"is_coffee_guess_p2": ["", "", "", "True", "True", "True", "True", "False"],
                      "has_description": [False, True, False, True, True, False, False, False]})
    s = labels.sample_roaster(g, 4, seed=1)
    assert len(s) == 4
    assert (s.is_coffee_guess_p2 == "").sum() == 2   # half the quota


# ---------- Excel round trip ----------

def small_csv(tmp_path):
    """A 2-product labels CSV with note lines, like the real one."""
    cols = ["label_id"] + labels.TEXT_COLUMNS + list(labels.FIELDS) + ["labeller_notes"]
    rows = [["1", "Kapi Kottai", "7311867969698", "Nātakurinji", "", "Coffee", "250 g / Whole Beans",
             "", "Washed arabica from Chikmagalur"] + [""] * 11,
            ["2", "Subko", "43461420023962", "Lot #SH3", "", "", "Whole Bean", "", ""] + [""] * 11]
    path = tmp_path / "hand_labels.csv"
    with path.open("w", encoding="utf-8-sig", newline="") as f:
        f.write("# note one\n# note two\n")
        pd.DataFrame(rows, columns=cols).to_csv(f, index=False)
    return path


GOOD = {"is_coffee": "yes", "origin_country": "India", "region": "Chikmagalur", "estate": "unknown",
        "species": "arabica", "origin_type": "single_origin", "process": "washed",
        "roast_level": "unknown", "decaf": "no", "flavoured_infused": "no"}


def fill(xlsx, values_by_row, product_id_override=None):
    """Type values into the Labels sheet, as a person would in Excel."""
    from openpyxl import load_workbook
    wb = load_workbook(xlsx)
    ws = wb["Labels"]
    header = [c.value for c in ws[1]]
    for i, values in enumerate(values_by_row, start=2):
        for field, value in values.items():
            ws.cell(row=i, column=header.index(field) + 1, value=value)
    if product_id_override is not None:
        ws.cell(row=2, column=header.index("product_id") + 1, value=product_id_override)
    wb.save(xlsx)


def test_xlsx_keeps_product_id_as_text_and_adds_dropdowns(tmp_path):
    from openpyxl import load_workbook
    csv, xlsx = small_csv(tmp_path), tmp_path / "hand_labels.xlsx"
    before = csv.read_bytes()
    labels.create_xlsx(csv, xlsx)
    assert csv.read_bytes() == before                          # CSV untouched
    wb = load_workbook(xlsx)
    assert wb.sheetnames == ["Instructions", "Labels"]
    ws = wb["Labels"]
    header = [c.value for c in ws[1]]
    pid = ws.cell(row=2, column=header.index("product_id") + 1)
    assert pid.value == "7311867969698" and pid.number_format == "@"
    assert ws.freeze_panes == labels.FROZEN
    dvs = ws.data_validations.dataValidation
    lists = {dv.formula1 for dv in dvs}
    assert '"yes,no,unknown"' in lists and '"arabica,robusta,blend,other,unknown"' in lists
    # One dropdown per list field (7); origin_country, region and estate stay free text.
    assert len(dvs) == sum(1 for a in labels.ALLOWED.values() if a) == 7
    assert all(dv.showErrorMessage for dv in dvs)              # invalid typing is rejected


def test_import_clean_workbook_writes_csv_and_keeps_notes(tmp_path):
    csv, xlsx = small_csv(tmp_path), tmp_path / "hand_labels.xlsx"
    labels.create_xlsx(csv, xlsx)
    fill(xlsx, [GOOD, {**GOOD, "species": "Other", "region": "unknown"}])   # capital O is fine
    assert labels.import_xlsx(csv, xlsx) == []
    text = csv.read_text(encoding="utf-8-sig")
    assert text.startswith("# note one\n# note two\n")
    out = labels.read_labels(csv)
    assert out.product_id.tolist() == ["7311867969698", "43461420023962"]
    assert out.species.tolist() == ["arabica", "other"]
    assert (tmp_path / "hand_labels.csv.bak").exists()


def test_import_reports_problems_and_leaves_csv_alone(tmp_path):
    csv, xlsx = small_csv(tmp_path), tmp_path / "hand_labels.xlsx"
    labels.create_xlsx(csv, xlsx)
    before = csv.read_bytes()
    # Row 1: product ID corrupted into a number, as Excel would; row 2: a bad value and a blank.
    fill(xlsx, [GOOD, {**GOOD, "process": "fermented", "decaf": None}],
         product_id_override=7.31187e12)
    problems = labels.import_xlsx(csv, xlsx)
    assert any("product_id" in p and "7311870000000" in p for p in problems)
    assert any("process = 'fermented'" in p for p in problems)
    assert any("decaf is blank" in p for p in problems)
    assert csv.read_bytes() == before


def test_india_rule_fills_unknown_only_and_reports_conflicts():
    df = pd.DataFrame({"label_id": ["1", "2", "3", "4", "5"],
                       "region": ["Chikmagalur", "Coorg", "Chiapas", "unknown", "Sakleshpur"],
                       "origin_country": ["unknown", "India", "Mexico", "unknown", "Brazil"]})
    change, conflicts = labels.india_rule_changes(df)
    assert change == [0]                                  # only the unknown one
    assert len(conflicts) == 1 and "Brazil" in conflicts[0]


def test_create_xlsx_never_overwrites(tmp_path):
    import pytest
    csv, xlsx = small_csv(tmp_path), tmp_path / "hand_labels.xlsx"
    labels.create_xlsx(csv, xlsx)
    with pytest.raises(SystemExit):
        labels.create_xlsx(csv, xlsx)


def test_read_labels_skips_note_lines(tmp_path, monkeypatch):
    f = tmp_path / "hand_labels.csv"
    f.write_text('# note one\n# note two, with "quotes"\nlabel_id,product_title\n1,Lot #08\n',
                 encoding="utf-8-sig")
    monkeypatch.setattr(labels, "LABELS", f)
    df = labels.read_labels()
    assert df.columns.tolist() == ["label_id", "product_title"]
    assert df.product_title.iloc[0] == "Lot #08"   # '#' inside a value is kept
